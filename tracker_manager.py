import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows

def load_or_create_tracker(tracker_path):
    headers = [
        "Invoice Firm Id", 
        "Finalized Date", 
        "Grand Total (Submitted)", 
        "Grand Total (Reduction)", 
        "Grand Total (Payable)", 
        "Link"
    ]
    
    if not os.path.exists(tracker_path):
        df = pd.DataFrame(columns=headers)
        df.to_excel(tracker_path, index=False)
        return df
    else:
        try:
            return pd.read_excel(tracker_path)
        except Exception as e:
            print(f"Error reading tracker: {e}")
            return pd.DataFrame(columns=headers)

def update_tracker(tracker_path, metadata, pdf_path):
    """
    Appends a new row to the tracker if the invoice is not already present.
    """
    df = load_or_create_tracker(tracker_path)
    
    invoice_id = metadata.get("invoice_number", "")
    finalized_date = metadata.get("finalized_date", "")
    submitted = metadata.get("total_submitted", 0.0)
    reduction = metadata.get("total_reduction", 0.0)
    payable = metadata.get("total_payable", 0.0)
    
    # Check for duplicate
    if invoice_id and str(invoice_id) in df["Invoice Firm Id"].astype(str).values:
        print(f"Invoice {invoice_id} already in tracker. Skipping.")
        return

    # Create Link Formula
    # Formula: =HYPERLINK("path", "Link")
    # We use the absolute path.
    link_formula = f'=HYPERLINK("{pdf_path}", "Link")'
    
    new_row = {
        "Invoice Firm Id": invoice_id,
        "Finalized Date": finalized_date,
        "Grand Total (Submitted)": submitted,
        "Grand Total (Reduction)": reduction,
        "Grand Total (Payable)": payable,
        "Link": link_formula
    }
    
    # Append using openpyxl to preserve formulas if any (though we are appending a formula)
    # Using pandas append is deprecated, use concat.
    
    # Actually, writing formulas with pandas is tricky. It usually writes as string.
    # To write formula as a working link, we might need openpyxl directly or use the XlsxWriter engine.
    # But for appending to an existing file, loading whole DF and saving might break existing formatting?
    # User said "tracker file wherein all new invoices shall be appended".
    
    # Let's use openpyxl to append strictly.
    
    wb = None
    if os.path.exists(tracker_path):
        wb = load_workbook(tracker_path)
        ws = wb.active
    else:
        # Should be created by load_or_create, but let's be safe.
        df = pd.DataFrame(columns=[
            "Invoice Firm Id", "Finalized Date", "Grand Total (Submitted)", 
            "Grand Total (Reduction)", "Grand Total (Payable)", "Link"
        ])
        df.to_excel(tracker_path, index=False)
        wb = load_workbook(tracker_path)
        ws = wb.active

    # Append row
    # Row sequence: ID, Date, Sub, Red, Pay, Link
    # Note: openpyxl rows are 1-based, columns 1-based.
    
    # Check for duplicates manually in openpyxl?
    # We already checked in DF properly.
    
    next_row = ws.max_row + 1
    
    ws.cell(row=next_row, column=1).value = invoice_id
    ws.cell(row=next_row, column=2).value = finalized_date
    ws.cell(row=next_row, column=3).value = submitted
    ws.cell(row=next_row, column=4).value = reduction
    ws.cell(row=next_row, column=5).value = payable
    ws.cell(row=next_row, column=6).value = link_formula
    
    wb.save(tracker_path)
    print(f"Updated tracker with Invoice {invoice_id}")
